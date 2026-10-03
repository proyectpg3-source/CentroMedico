from pathlib import Path
from zipfile import ZipFile

from docx import Document

docx_path = Path(__file__).resolve().parent.parent / "reporte-ux-login-centro-medico.docx"
doc = Document(docx_path)

headings = [
    paragraph.text
    for paragraph in doc.paragraphs
    if paragraph.style.name.startswith("Heading")
]

with ZipFile(docx_path) as archive:
    media_files = [name for name in archive.namelist() if name.startswith("word/media/")]

print(f"archivo={docx_path}")
print(f"parrafos={len(doc.paragraphs)}")
print(f"tablas={len(doc.tables)}")
print(f"imagenes={len(media_files)}")
print("encabezados=" + " | ".join(headings))
