# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


BASE_DIR = Path(__file__).resolve().parent
CAPTURE_DIR = BASE_DIR / "captures"
OUTPUT = BASE_DIR.parent / "reporte-ux-login-centro-medico.docx"
SKETCH = CAPTURE_DIR / "boceto-propuesta-login.png"


def font(size, bold=False):
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            continue
    return ImageFont.load_default()


def fit_text(draw, text, max_width, font_obj):
    words = text.split()
    lines = []
    current = ""
    for word in words:
        trial = word if not current else f"{current} {word}"
        width = draw.textbbox((0, 0), trial, font=font_obj)[2]
        if width <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_text_box(draw, xy, text, font_obj, fill, max_width, line_gap=6):
    x, y = xy
    for line in fit_text(draw, text, max_width, font_obj):
        draw.text((x, y), line, fill=fill, font=font_obj)
        y += font_obj.size + line_gap
    return y


def crop_capture(name, box):
    src = CAPTURE_DIR / name
    dst = CAPTURE_DIR / f"crop-{name}"
    img = Image.open(src).convert("RGB")
    img.crop(box).save(dst, quality=92)
    return dst


def create_sketch():
    CAPTURE_DIR.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (1200, 720), "#f3f7f5")
    draw = ImageDraw.Draw(img)

    draw.rounded_rectangle((44, 44, 1156, 676), radius=28, fill="#ffffff", outline="#d8e4df", width=3)
    draw.rectangle((44, 44, 510, 676), fill="#102436")
    draw.rounded_rectangle((44, 44, 510, 676), radius=28, outline="#102436", width=1)

    draw.rounded_rectangle((84, 84, 142, 142), radius=14, fill="#0f6e56")
    draw.text((102, 96), "+", fill="#ffffff", font=font(34, True))
    draw.text((162, 86), "Centro Medico", fill="#ffffff", font=font(23, True))
    draw.text((162, 116), "La Aurora", fill="#a9c8bd", font=font(17))

    draw_text_box(
        draw,
        (84, 224),
        "Acceso por area con instrucciones visibles y menos pasos antes de escribir credenciales.",
        font(30, True),
        "#ffffff",
        350,
        10,
    )
    draw_text_box(
        draw,
        (84, 390),
        "La seleccion de rol se muestra de inmediato y cada campo conserva su etiqueta aunque el usuario escriba.",
        font(17),
        "#b9c9d5",
        340,
        7,
    )

    card = (590, 86, 1118, 632)
    draw.rounded_rectangle(card, radius=22, fill="#ffffff", outline="#d8e4df", width=2)
    draw.text((636, 132), "Entrar al sistema", fill="#132333", font=font(30, True))
    draw.text((636, 174), "Selecciona tu area e ingresa tus datos.", fill="#65758a", font=font(17))

    roles = [
        ("Recepcion", "#f59e0b"),
        ("Enfermeria", "#10b981"),
        ("Medico", "#0ea5e9"),
        ("Admin", "#64748b"),
        ("Fundacion", "#06b6d4"),
    ]
    x = 636
    y = 224
    for index, (label, color) in enumerate(roles):
        bx = x + (index % 2) * 218
        by = y + (index // 2) * 58
        draw.rounded_rectangle((bx, by, bx + 196, by + 42), radius=10, fill="#f8fafc", outline="#d6dde6")
        draw.ellipse((bx + 12, by + 12, bx + 30, by + 30), fill=color)
        draw.text((bx + 42, by + 10), label, fill="#243447", font=font(16, True))

    field_y = 406
    for label, placeholder in [("Correo electronico", "usuario@centromedico.com"), ("Contrasena", "********")]:
        draw.text((636, field_y), label.upper(), fill="#556274", font=font(13, True))
        draw.rounded_rectangle((636, field_y + 22, 1074, field_y + 70), radius=12, fill="#f8fafc", outline="#cfd8e3")
        draw.text((656, field_y + 36), placeholder, fill="#8794a6", font=font(16))
        field_y += 92

    draw.rounded_rectangle((636, 590, 1074, 640), radius=14, fill="#0f6e56")
    draw.text((820, 604), "Ingresar", fill="#ffffff", font=font(18, True))
    img.save(SKETCH, quality=95)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_border(cell, color="D9E2EA"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)
        borders.append(tag)


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.allow_autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def set_cell_text(cell, text, bold=False, color=None):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(46, 116, 181) if level <= 2 else RGBColor(31, 77, 120)
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.10
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.10
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    return paragraph


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.italic = True
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(91, 103, 119)


def add_image_pair(doc, left_path, right_path, caption):
    table = doc.add_table(rows=2, cols=2)
    set_table_width(table, [3.05, 3.05])
    for idx, title in enumerate(("Antes", "Despues")):
        cell = table.cell(0, idx)
        set_cell_shading(cell, "F2F4F7")
        set_cell_border(cell)
        set_cell_text(cell, title, bold=True, color="1F4D78")
    for idx, img_path in enumerate((left_path, right_path)):
        cell = table.cell(1, idx)
        set_cell_border(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(img_path), width=Inches(2.95))
    add_caption(doc, caption)


def create_report():
    create_sketch()
    before_roles = crop_capture("antes-02-areas.png", (270, 40, 1095, 708))
    after_roles = crop_capture("despues-02-areas.png", (270, 40, 1095, 708))
    before_form = crop_capture("antes-03-formulario.png", (430, 150, 930, 620))
    after_form = crop_capture("despues-03-formulario.png", (430, 130, 930, 640))

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("Reporte de mejora UX - Modulo de Login")
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.bold = True
    run.font.color.rgb = RGBColor(15, 110, 86)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle_run = subtitle.add_run("Centro Medico La Aurora | 28 de agosto de 2026")
    subtitle_run.font.name = "Calibri"
    subtitle_run.font.size = Pt(11)
    subtitle_run.font.color.rgb = RGBColor(91, 103, 119)

    meta = doc.add_table(rows=3, cols=2)
    set_table_width(meta, [1.65, 4.85])
    meta_items = [
        ("Modulo seleccionado", "Inicio de sesion y seleccion de area/rol."),
        ("Objetivo", "Analizar la interfaz actual del login, detectar problemas de usabilidad y documentar una mejora minima implementada."),
        ("Principios evaluados", "Claridad, consistencia, visibilidad, eficiencia y retroalimentacion."),
    ]
    for row, (label, value) in zip(meta.rows, meta_items):
        set_cell_shading(row.cells[0], "E8EEF5")
        set_cell_border(row.cells[0])
        set_cell_border(row.cells[1])
        set_cell_text(row.cells[0], label, bold=True, color="1F4D78")
        set_cell_text(row.cells[1], value)

    add_heading(doc, "1. Analisis de la interfaz actual", 1)
    add_body(
        doc,
        "El login actual tiene una identidad visual atractiva y consistente con un sistema medico: utiliza fondo fotografico, iconografia clinica, colores institucionales y una navegacion por roles. La revision se hizo sobre el prototipo en desarrollo ubicado en la ruta /login.",
    )
    add_bullet(doc, "Claridad: la pantalla inicial comunica marca, pero el paso de seleccion de area no muestra de inmediato todas las opciones.")
    add_bullet(doc, "Consistencia: los iconos y colores se mantienen entre pantalla de bienvenida, roles y formulario.")
    add_bullet(doc, "Visibilidad: algunos controles importantes estaban ocultos hasta hacer clic en el boton central.")
    add_bullet(doc, "Eficiencia: el usuario necesitaba un clic adicional para descubrir las areas disponibles.")
    add_bullet(doc, "Retroalimentacion: el formulario muestra estados de carga y mensajes de error al iniciar sesion.")

    add_heading(doc, "2. Problemas detectados", 1)
    problems = doc.add_table(rows=1, cols=4)
    set_table_width(problems, [1.15, 1.6, 1.85, 1.9])
    headers = ["Problema", "Principio afectado", "Por que afecta", "Mejora propuesta"]
    for idx, header in enumerate(headers):
        cell = problems.cell(0, idx)
        set_cell_shading(cell, "F2F4F7")
        set_cell_border(cell)
        set_cell_text(cell, header, bold=True, color="1F4D78")
    rows = [
        (
            "Areas ocultas",
            "Visibilidad y eficiencia",
            "El usuario no ve las opciones de rol al entrar al segundo paso; debe descubrir que el boton central despliega el menu.",
            "Mostrar las areas automaticamente y agregar una instruccion breve debajo del menu.",
        ),
        (
            "Etiquetas no visibles",
            "Claridad y accesibilidad",
            "Los campos dependian del placeholder; al escribir se pierde la pista visual de que dato corresponde a cada campo.",
            "Mantener etiquetas visibles sobre correo electronico y contrasena.",
        ),
    ]
    for row_data in rows:
        cells = problems.add_row().cells
        for idx, value in enumerate(row_data):
            set_cell_border(cells[idx])
            set_cell_text(cells[idx], value)

    doc.add_page_break()

    add_heading(doc, "3. Propuesta de mejora UX", 1)
    add_body(
        doc,
        "La propuesta conserva el estilo visual que ya existia, pero reduce la friccion del flujo. En lugar de esconder la seleccion de area, el sistema muestra de inmediato las opciones disponibles. Tambien mantiene visibles las etiquetas de los campos para que el usuario no dependa solo del texto placeholder.",
    )
    add_bullet(doc, "Accion 1: mostrar Recepcion, Enfermeria, Medico, Administrador y Fundacion apenas el usuario entra al paso de areas.")
    add_bullet(doc, "Accion 2: agregar una frase de orientacion: 'Elige tu area para continuar con tus credenciales'.")
    add_bullet(doc, "Accion 3: mostrar etiquetas persistentes en los campos de correo electronico y contrasena.")

    doc.add_picture(str(SKETCH), width=Inches(6.2))
    add_caption(doc, "Boceto rapido de la propuesta: roles visibles, instrucciones breves y campos con etiquetas.")

    add_heading(doc, "4. Implementacion minima", 1)
    add_body(
        doc,
        "Se implementaron dos mejoras visuales pequenas en el modulo de login. La primera cambia el comportamiento de entrada al paso de areas para que el menu se despliegue automaticamente. La segunda agrega etiquetas visibles a los campos del formulario. La validacion funcional por rol se mantiene: el backend rechaza credenciales si la cuenta no corresponde al area seleccionada.",
    )
    add_bullet(doc, "Archivo principal modificado: frontend/src/app/pages/Login.tsx.")
    add_bullet(doc, "Verificacion tecnica: el frontend compilo correctamente con npm run build y el backend paso npm test.")

    doc.add_page_break()

    add_heading(doc, "5. Capturas antes y despues", 1)
    add_body(
        doc,
        "Las siguientes capturas documentan el cambio implementado. La comparacion se centra en visibilidad y claridad, que fueron los dos problemas detectados en la revision.",
    )
    add_image_pair(
        doc,
        before_roles,
        after_roles,
        "Comparacion 1: antes las areas estaban ocultas; despues se muestran automaticamente.",
    )
    add_image_pair(
        doc,
        before_form,
        after_form,
        "Comparacion 2: antes los campos dependian del placeholder; despues tienen etiquetas visibles.",
    )

    add_heading(doc, "Conclusion", 1)
    add_body(
        doc,
        "La mejora reduce pasos innecesarios y aumenta la claridad sin reemplazar la identidad visual del login. Para una entrega academica, el modulo demuestra analisis de usabilidad, propuesta practica, boceto e implementacion minima con evidencia antes/despues.",
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("Reporte UX - Centro Medico La Aurora")
    footer_run.font.name = "Calibri"
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(91, 103, 119)

    doc.save(OUTPUT)


if __name__ == "__main__":
    create_report()
    print(OUTPUT)
